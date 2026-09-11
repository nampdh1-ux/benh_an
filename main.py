import io
import re
import base64
from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.responses import StreamingResponse
import base64
import hashlib
import html
import io
import json
import os
import random
import re
import smtplib
import unicodedata
import urllib.parse
import urllib.request
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from typing import Any, Dict, List
from fastapi.staticfiles import StaticFiles
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, Response
from fastapi.templating import Jinja2Templates
from fpdf import FPDF
from google import genai
from google.genai import types
from PIL import Image
from pydantic import BaseModel

app = FastAPI(title="Bệnh Án Lâm Sàng Win2K")
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

# Cấu hình môi trường & AI
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
MODEL_DEFAULT = "gemini-3.1-flash-lite"

# Font Unicode cho FPDF
FONT_REGULAR = "Roboto-Regular.ttf"
FONT_BOLD = "Roboto-Bold.ttf"

def download_fonts_if_missing():
    urls = {
        FONT_REGULAR: "https://raw.githubusercontent.com/google/fonts/main/apache/roboto/Roboto-Regular.ttf",
        FONT_BOLD: "https://raw.githubusercontent.com/google/fonts/main/apache/roboto/Roboto-Bold.ttf"
    }
    for filename, url in urls.items():
        if not os.path.exists(filename) or os.path.getsize(filename) < 10000:
            try:
                req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
                with urllib.request.urlopen(req, timeout=15) as resp, open(filename, 'wb') as f:
                    f.write(resp.read())
            except Exception as e:
                print(f"Chưa thể tải {filename}: {e}")

download_fonts_if_missing()

def get_ai_client():
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key:
        return None
    try:
        return genai.Client(api_key=api_key)
    except Exception as e:
        print(f"Lỗi khởi tạo AI Client: {e}")
        return None

def strip_accents(text: Any) -> str:
    if not text:
        return ""
    text = unicodedata.normalize('NFD', str(text))
    text = re.sub(r'[\u0300-\u036f]', '', text)
    return text.replace('đ', 'd').replace('Đ', 'D')

def format_bullet_points(text: str) -> str:
    if not text or not str(text).strip():
        return "Chưa ghi nhận thông tin."
    lines = str(text).strip().split("\n")
    formatted_lines = []
    for line in lines:
        cleaned = line.strip()
        if cleaned:
            if not cleaned.startswith("-") and not cleaned.startswith("*"):
                formatted_lines.append(f"- {cleaned}")
            else:
                formatted_lines.append(cleaned)
    return "\n".join(formatted_lines)

def get_benh_su_text(payload: Dict[str, Any]) -> str:
    if payload.get("loai_benh_an") == "Hậu phẫu":
        return f"- Trước mổ: {payload.get('bs_truoc_mo', '')}\n- Trong mổ: {payload.get('bs_trong_mo', '')}\n- Sau mổ: {payload.get('bs_sau_mo', '')}"
    return payload.get("benh_su", "")

def get_prioritized_organs(payload: Dict[str, Any]):
    organs = [
        ("Tuần hoàn", "kham_tuan_hoan"),
        ("Hô hấp", "kham_ho_hap"),
        ("Tiêu hóa", "kham_tieu_hoa"),
        ("Thần kinh", "kham_than_kinh"),
        ("Thận - Tiết niệu", "kham_tiet_nieu"),
        ("Cơ xương khớp", "kham_co_xuong_khop"),
        ("Các cơ quan khác", "kham_co_quan_khac")
    ]
    favored_key = payload.get("uu_tien_co_quan", "none")
    if favored_key != "none":
        organs = [organ for organ in organs if organ[1] == favored_key] + [
            organ for organ in organs if organ[1] != favored_key
        ]
    return organs

@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    return templates.TemplateResponse(request, "index.html")

# --- CÁC ENDPOINT AI ---
@app.post("/api/ai/cdpb")
async def api_ai_cdpb(payload: Dict[str, Any]):
    client = get_ai_client()
    if not client:
        raise HTTPException(status_code=500, detail="Chưa cấu hình GEMINI_API_KEY!")
    
    context = payload.get("full_context", "")
    prompt = f"""
Bạn là một bác sĩ chuyên khoa thực thụ. Dưới đây là toàn bộ dữ liệu lâm sàng thu thập được từ đầu đến thời điểm thăm khám hiện tại:
==================================================
{context}
==================================================

Dựa trên nguyên lý biện luận lâm sàng (Clinical Reasoning):
1. Đưa ra danh sách các Chẩn đoán phân biệt (Differential Diagnoses), sắp xếp theo thứ tự ưu tiên đúng hoặc mức độ nguy cấp, dạng
    1. A
    2. B
    3. C
    ... Không giải thích hay câu từ gì thêm
2. Viết đoạn Biện luận chẩn đoán sơ bộ: Phân tích logic tại sao hướng tới chẩn đoán sơ bộ (dấu hiệu chỉ điểm, yếu tố nguy cơ) và tại sao chưa thể loại trừ các chẩn đoán phân biệt, vào thẳng vấn đề, dưới dạng xuống dòng đơn giản.

YÊU CẦU ĐỊNH DẠNG: Trả về ĐÚNG 2 thẻ:
[CHAN_DOAN_PHAN_BIET]
...
[BIEN_LUAN_SO_BO]
...
"""
    try:
        response = client.models.generate_content(model=MODEL_DEFAULT, contents=prompt)
        resp = response.text or ""
        cdpb, bl = "", ""
        if "[CHAN_DOAN_PHAN_BIET]" in resp and "[BIEN_LUAN_SO_BO]" in resp:
            parts = resp.split("[BIEN_LUAN_SO_BO]")
            cdpb = parts[0].replace("[CHAN_DOAN_PHAN_BIET]", "").strip()
            bl = parts[1].strip()
        else:
            cdpb = resp.strip()
        return {"chan_doan_phan_biet": cdpb, "bien_luan": bl}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi AI: {str(e)}")

@app.post("/api/ai/treatment")
async def api_ai_treatment(payload: Dict[str, Any]):
    client = get_ai_client()
    if not client:
        raise HTTPException(status_code=500, detail="Chưa cấu hình GEMINI_API_KEY!")
    
    context = payload.get("full_context", "")
    prompt = f"""
Bạn là bác sĩ điều trị. Dưới đây là toàn bộ hồ sơ bệnh nhân tính đến khi đã có Chẩn đoán xác định và Cận lâm sàng:
==================================================
{context}
==================================================

Hãy xây dựng kế hoạch điều trị toàn diện theo y học thực chứng:
1. Mục tiêu điều trị.
2. Điều trị cụ thể: Bao gồm chế độ chăm sóc/dinh dưỡng, dùng thuốc (tên hoạt chất, liều lượng, đường dùng nếu cần thiết) hoặc can thiệp ngoại khoa/chăm sóc hậu phẫu chuyên biệt.
3. Kế hoạch theo dõi: Các dấu hiệu sinh tồn, dẫn lưu, biến chứng cần tầm soát.
    Vào thẳng vấn đề, dưới dạng xuống dòng, đơn giản, không màu mè.
YÊU CẦU ĐỊNH DẠNG: Trả về ĐÚNG 3 thẻ:
[MUC_TIEU]
...
[DIEU_TRI_CU_THE]
...
[THEO_DOI]
...
"""
    try:
        response = client.models.generate_content(model=MODEL_DEFAULT, contents=prompt)
        txt = response.text or ""
        mt, ct, td = "", "", ""
        if "[MUC_TIEU]" in txt and "[DIEU_TRI_CU_THE]" in txt and "[THEO_DOI]" in txt:
            p1 = txt.split("[DIEU_TRI_CU_THE]")
            mt = p1[0].replace("[MUC_TIEU]", "").strip()
            p2 = p1[1].split("[THEO_DOI]")
            ct = p2[0].strip()
            td = p2[1].strip()
        return {"dt_muc_tieu": mt, "dt_cu_the": ct, "dt_theo_doi": td}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi AI: {str(e)}")

@app.post("/api/ai/prognosis")
async def api_ai_prognosis(payload: Dict[str, Any]):
    client = get_ai_client()
    if not client:
        raise HTTPException(status_code=500, detail="Chưa cấu hình GEMINI_API_KEY!")
    
    context = payload.get("full_context", "")
    prompt = f"""
Bạn là bác sĩ lâm sàng. Dưới đây là toàn bộ diễn biến ca bệnh và phương án điều trị đã thiết lập:
==================================================
{context}
==================================================

Hãy phân tích:
1. Tiên lượng: Gồm tiên lượng gần (biến chứng cấp, khả năng hồi phục trong đợt điều trị) và tiên lượng xa (tái phát, di chứng, chức năng cơ quan).
2. Tư vấn & Giáo dục sức khỏe: Hướng dẫn chăm sóc, chế độ vận động/ăn uống, dấu hiệu báo động đỏ cần tái khám ngay.
    Vào thẳng vấn đề, dưới dạng xuống dòng, đơn giản, không màu mè.
YÊU CẦU ĐỊNH DẠNG: Trả về ĐÚNG 2 thẻ:
[TIEN_LUONG]
...
[TU_VAN]
...
"""
    try:
        response = client.models.generate_content(model=MODEL_DEFAULT, contents=prompt)
        res_text = response.text or ""
        tl, tv = "", ""
        if "[TIEN_LUONG]" in res_text and "[TU_VAN]" in res_text:
            parts = res_text.split("[TU_VAN]")
            tl = parts[0].replace("[TIEN_LUONG]", "").strip()
            tv = parts[1].strip()
        return {"tien_luong": tl, "tu_van": tv}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi AI: {str(e)}")

@app.post("/api/ai/critique")
async def api_ai_critique(payload: Dict[str, Any]):
    client = get_ai_client()
    if not client:
        raise HTTPException(status_code=500, detail="Chưa cấu hình GEMINI_API_KEY!")
    phong_cach = payload.get("phong_cach", "Học thuật & Hướng dẫn")
    context = f"Bệnh sử: {get_benh_su_text(payload)}\nChẩn đoán SB: {payload.get('chan_doan_so_bo')}\nChẩn đoán XĐ: {payload.get('chan_doan_xac_dinh')}"
    prompt = f"""Bạn là Giảng viên lâm sàng. Nhận xét ca bệnh ({context}) theo phong cách {phong_cach}.
    Trả về ĐÚNG định dạng JSON thuần: {{"nhan_xet_tong_the": "...", "danh_sach_cau_hoi": [{{"chu_de": "...", "cau_hoi": "...", "goi_y_tra_loi": "..."}}]}}"""
    try:
        response = client.models.generate_content(
            model=MODEL_DEFAULT,
            contents=prompt,
        )
        res_pb = (response.text or "").strip()
        if res_pb.startswith("```json"): res_pb = res_pb[7:]
        if res_pb.startswith("```"): res_pb = res_pb[3:]
        if res_pb.endswith("```"): res_pb = res_pb[:-3]
        return json.loads(res_pb.strip())
    except Exception as e:
        return {"nhan_xet_tong_the": f"Lỗi phản biện: {str(e)}", "danh_sach_cau_hoi": []}

@app.post("/api/ocr/batch")
async def api_ocr_batch(
    files: List[UploadFile] = File(...),
    context: str = Form("")
):
    client = get_ai_client()
    if not client:
        raise HTTPException(status_code=500, detail="Chưa cấu hình GEMINI_API_KEY!")
    
    clinical_ctx_str = "Chưa có thông tin ngữ cảnh lâm sàng."
    if context:
        try:
            ctx_data = json.loads(context)
            loai_ba = ctx_data.get("loai_benh_an", "Nội khoa / Tiền phẫu")
            spo2_val = ctx_data.get('sh_spo2', '')
            spo2_str = f" | SpO2: {spo2_val}%" if spo2_val else ""
            vital_str = f"Mạch: {ctx_data.get('sh_mach', '--')} ck/p | HA: {ctx_data.get('sh_ha', '--')} mmHg | Thở: {ctx_data.get('sh_nhip_tho', '--')} l/p{spo2_str}"

            if loai_ba == "Hậu phẫu":
                clinical_ctx_str = f"Loại: HẬU PHẪU\nBệnh nhân: {ctx_data.get('ho_ten')} ({ctx_data.get('tuoi')}t, {ctx_data.get('gioi_tinh')})\nLý do: {ctx_data.get('ly_do_vao_vien')}\nTrước mổ: {ctx_data.get('bs_truoc_mo')}\nTrong mổ: {ctx_data.get('bs_trong_mo')}\nSau mổ: {ctx_data.get('bs_sau_mo')}\nNgày HP: {ctx_data.get('ngay_hau_phau')}\nSinh hiệu: {vital_str}\nVết mổ: {ctx_data.get('kham_vet_mo')}\nDẫn lưu: {ctx_data.get('kham_dan_luu')}\nCĐ Sơ bộ: {ctx_data.get('chan_doan_so_bo')}"
            else:
                clinical_ctx_str = f"Loại: NỘI KHOA\nBệnh nhân: {ctx_data.get('ho_ten')} ({ctx_data.get('tuoi')}t, {ctx_data.get('gioi_tinh')})\nLý do: {ctx_data.get('ly_do_vao_vien')}\nBệnh sử: {ctx_data.get('benh_su')}\nTiền sử: {ctx_data.get('ts_noi_khoa')}\nSinh hiệu: {vital_str}\nCĐ Sơ bộ: {ctx_data.get('chan_doan_so_bo')}"
        except Exception:
            clinical_ctx_str = str(context)

    ocr_prompt = f"""Bạn là bác sĩ lâm sàng. Đọc cận lâm sàng đính kèm dựa trên ngữ cảnh, không câu dẫn, vào thẳng vấn đề:\n{clinical_ctx_str}\nTrả về JSON có 2 key: 'ket_qua' (chỉ số đo được) và 'phien_giai' (biện luận theo bệnh cảnh)."""
    results = []
    for file in files:
        try:
            raw_bytes = await file.read()
            if not raw_bytes: continue
            
            img = Image.open(io.BytesIO(raw_bytes))
            try:
                import PIL.ImageOps as ImageOps
                img = ImageOps.exif_transpose(img)
            except Exception: pass
            
            if img.mode != 'RGB': img = img.convert('RGB')
            if max(img.size) > 1600: img.thumbnail((1600, 1600), Image.Resampling.LANCZOS)
            
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=85)
            
            image_part = types.Part.from_bytes(data=buf.getvalue(), mime_type="image/jpeg")
            response = client.models.generate_content(
                model=MODEL_DEFAULT,
                contents=[image_part, ocr_prompt],
                config=types.GenerateContentConfig(response_mime_type="application/json")
            )
            resp_text = (response.text or "").strip()
            if resp_text.startswith("```json"): resp_text = resp_text[7:]
            if resp_text.startswith("```"): resp_text = resp_text[3:]
            if resp_text.endswith("```"): resp_text = resp_text[:-3]
            parsed = json.loads(resp_text.strip())
            results.append({"ket_qua": parsed.get("ket_qua", "Không đọc được chỉ số."), "phien_giai": parsed.get("phien_giai", "-")})
        except Exception as err:
            results.append({"ket_qua": f"Lỗi xử lý: {str(err)}", "phien_giai": "-"})
    return {"results": results}

# --- PDF ENGINE HỖ TRỢ NHIỀU ẢNH TRÊN MỖI DÒNG CLS ---
class RobustUnicodePDF(FPDF):
    def __init__(self, loai_ba="Nội khoa / Tiền phẫu", *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.loai_ba = loai_ba
        self.use_unicode = False
        if os.path.exists(FONT_REGULAR) and os.path.exists(FONT_BOLD):
            try:
                self.add_font("Roboto", "", FONT_REGULAR)
                self.add_font("Roboto", "B", FONT_BOLD)
                self.font_family_name = "Roboto"
                self.use_unicode = True
            except Exception:
                self.font_family_name = "Helvetica"
        else:
            self.font_family_name = "Helvetica"

    def clean_text(self, text: Any) -> str:
        s = str(text or "")
        return s if self.use_unicode else strip_accents(s)

    def header(self):
        if self.page_no() == 1:
            self.set_font(self.font_family_name, "B" if not self.use_unicode else "", 15)
            tieu_de = "BỆNH ÁN HẬU PHẪU" if self.loai_ba == "Hậu phẫu" else "BỆNH ÁN LÂM SÀNG"
            self.cell(0, 8, self.clean_text(tieu_de), align="C", new_x="LMARGIN", new_y="NEXT")
            self.set_font(self.font_family_name, "", 9)
            self.cell(0, 4, self.clean_text(f"Thời gian lập: {datetime.now().strftime('%d/%m/%Y %H:%M')}"), align="C", new_x="LMARGIN", new_y="NEXT")
            self.ln(4)

    def add_sec(self, title: str):
        self.set_font(self.font_family_name, "B" if not self.use_unicode else "", 11)
        self.set_fill_color(225, 235, 245)
        self.cell(0, 7, self.clean_text(title), fill=True, new_x="LMARGIN", new_y="NEXT")
        self.ln(1)

    def add_subsec(self, title: str):
        self.set_font(self.font_family_name, "B", 10)
        self.cell(0, 6, self.clean_text(title), new_x="LMARGIN", new_y="NEXT")

    def add_txt(self, text: str):
        self.set_font(self.font_family_name, "", 9.5)
        self.multi_cell(0, 5, self.clean_text(text) if str(text).strip() else self.clean_text("Chưa ghi nhận thông tin."))
        self.ln(2)

    def render_table_cls(self, cls_rows):
        col_w = (self.w - self.l_margin - self.r_margin) / 2.0
        line_h = 5.0
        self.set_font(self.font_family_name, "B" if not self.use_unicode else "", 9.5)
        self.set_fill_color(230, 235, 245)
        if self.get_y() > 230:
            self.add_page()
        self.cell(col_w, 7, self.clean_text("KẾT QUẢ CẬN LÂM SÀNG & HÌNH ẢNH"), border=1, align="C", fill=True)
        self.cell(col_w, 7, self.clean_text("PHIÊN GIẢI / BIỆN GIẢI KẾT QUẢ"), border=1, align="C", fill=True, new_x="LMARGIN", new_y="NEXT")
        
        self.set_font(self.font_family_name, "", 9)
        for kq, pg, img_list in cls_rows:
            txt_kq = format_bullet_points(kq) if kq else "-"
            txt_pg = format_bullet_points(pg) if pg else "-"
            
            pil_images = []
            single_img_h = 42.0
            
            if isinstance(img_list, list):
                for img_b64 in img_list:
                    if img_b64 and len(str(img_b64).strip()) > 50:
                        try:
                            clean_b64 = str(img_b64)
                            if "," in clean_b64:
                                clean_b64 = clean_b64.split(",", 1)[1]
                            img_data = base64.b64decode(clean_b64)
                            im = Image.open(io.BytesIO(img_data))
                            if im.mode not in ("RGB", "L"):
                                im = im.convert("RGB")
                            pil_images.append(im)
                        except Exception as err:
                            txt_kq += f"\n[Lỗi ảnh: {str(err)}]"

            total_images_h = len(pil_images) * (single_img_h + 3)
            nb_l = len(self.multi_cell(col_w - 4, line_h, self.clean_text(txt_kq), dry_run=True, output="LINES"))
            nb_r = len(self.multi_cell(col_w - 4, line_h, self.clean_text(txt_pg), dry_run=True, output="LINES"))
            
            col_l_h = nb_l * line_h + total_images_h
            col_r_h = nb_r * line_h
            row_h = max(max(col_l_h, col_r_h) + 6, 12)
            
            if self.get_y() + row_h > 270:
                self.add_page()
                self.set_font(self.font_family_name, "B" if not self.use_unicode else "", 9.5)
                self.set_fill_color(230, 235, 245)
                self.cell(col_w, 7, self.clean_text("KẾT QUẢ CẬN LÂM SÀNG & HÌNH ẢNH"), border=1, align="C", fill=True)
                self.cell(col_w, 7, self.clean_text("PHIÊN GIẢI / BIỆN GIẢI KẾT QUẢ"), border=1, align="C", fill=True, new_x="LMARGIN", new_y="NEXT")
                self.set_font(self.font_family_name, "", 9)

            curr_x = self.get_x()
            curr_y = self.get_y()
            self.rect(curr_x, curr_y, col_w, row_h)
            self.rect(curr_x + col_w, curr_y, col_w, row_h)

            self.set_xy(curr_x + 2, curr_y + 2)
            self.multi_cell(col_w - 4, line_h, self.clean_text(txt_kq))
            
            curr_img_y = self.get_y() + 2
            for im in pil_images:
                try:
                    self.image(im, x=curr_x + 3, y=curr_img_y, w=col_w - 6, h=single_img_h)
                    curr_img_y += single_img_h + 3
                except Exception as img_err:
                    self.set_xy(curr_x + 2, curr_img_y)
                    self.multi_cell(col_w - 4, line_h, self.clean_text(f"[Lỗi vẽ ảnh: {str(img_err)}]"))
                    curr_img_y += line_h * 2

            self.set_xy(curr_x + col_w + 2, curr_y + 2)
            self.multi_cell(col_w - 4, line_h, self.clean_text(txt_pg))
            self.set_xy(curr_x, curr_y + row_h)
        self.ln(3)

@app.post("/api/export/pdf")
async def api_export_pdf(payload: Dict[str, Any]):
    try:
        if not os.path.exists(FONT_REGULAR):
            download_fonts_if_missing()

        loai_ba = str(payload.get("loai_benh_an", "Nội khoa / Tiền phẫu")).strip()
        is_hau_phau = (loai_ba == "Hậu phẫu")
        
        pdf = RobustUnicodePDF(loai_ba=loai_ba)
        pdf.add_page()

        pdf.add_sec("I. PHẦN HÀNH CHÍNH")
        hc = (
            f"- Họ và tên: {str(payload.get('ho_ten', '')).upper()}   |   Tuổi: {payload.get('tuoi')}   |   Giới tính: {payload.get('gioi_tinh')}\n"
            f"- Dân tộc: {payload.get('dan_tok')}   |   Nghề nghiệp: {payload.get('nghe_nghiep')}\n"
            f"- Khoa phòng: {payload.get('khoa_phong')}   |   Địa chỉ: {payload.get('dia_chi', '')}\n"
            f"- Ngày giờ vào viện: {payload.get('ngay_vao_vien')}   |   Ngày làm BA: {payload.get('ngay_lam_benh_an', '')}\n"
            f"- Người làm bệnh án: {payload.get('sinh_vien')}"
        )
        pdf.add_txt(hc)

        pdf.add_sec("II. LÝ DO VÀO VIỆN")
        pdf.add_txt(payload.get("ly_do_vao_vien", ""))

        pdf.add_sec("III. BỆNH SỬ")
        if is_hau_phau:
            pdf.add_subsec("1. Tình trạng trước mổ:")
            pdf.add_txt(format_bullet_points(payload.get("bs_truoc_mo", "")))
            pdf.add_subsec("2. Tình trạng trong mổ:")
            pdf.add_txt(format_bullet_points(payload.get("bs_trong_mo", "")))
            pdf.add_subsec("3. Quá trình sau mổ:")
            pdf.add_txt(format_bullet_points(payload.get("bs_sau_mo", "")))
        else:
            pdf.add_txt(payload.get("benh_su", ""))

        pdf.add_sec("IV. TIỀN SỬ")
        pdf.add_subsec("1. Tiền sử nội khoa:")
        pdf.add_txt(format_bullet_points(payload.get("ts_noi_khoa", "")))
        pdf.add_subsec("2. Tiền sử ngoại khoa & Dị ứng:")
        pdf.add_txt(format_bullet_points(payload.get("ts_ngoai_khoa", "")))
        pdf.add_subsec("3. Tiền sử bản thân (Lối sống & Thói quen):")
        pdf.add_txt(format_bullet_points(payload.get("ts_loi_song", "")))
        pdf.add_subsec("4. Tiền sử gia đình:")
        pdf.add_txt(format_bullet_points(payload.get("ts_gia_dinh", "")))

        pdf.add_sec("V. THĂM KHÁM LÂM SÀNG")
        if not is_hau_phau:
            pdf.add_subsec("1. Thăm khám lúc vào viện:")
            pdf.add_txt(format_bullet_points(payload.get("kham_vao_vien", "")))
            pdf.add_subsec("2. Thăm khám hiện tại - Toàn thân:")
        else:
            pdf.add_subsec(f"1. Thăm khám hiện tại ({payload.get('ngay_hau_phau', 'Hậu phẫu')}):")
            pdf.add_subsec("a. Toàn thân:")
            
        pdf.add_txt(format_bullet_points(payload.get("kham_toan_than", "")))
        
        mach = payload.get("sh_mach") or "--"
        nhiet = payload.get("sh_nhiet_do") or "--"
        ha = payload.get("sh_ha") or "--"
        nt = payload.get("sh_nhip_tho") or "--"
        spo2 = payload.get("sh_spo2") or "--"
        cn = payload.get("sh_can_nang") or "--"
        cc = payload.get("sh_chieu_cao") or "--"
        bmi = payload.get("sh_bmi") or "--"
        eval_bmi = payload.get("sh_bmi_eval") or "--"
        sh_line = f"Sinh hiệu: Mạch: {mach} ck/phút | HA: {ha} mmHg | Nhiệt độ: {nhiet} °C | Nhịp thở: {nt} l/phút | SpO2: {spo2}%\nThể trạng: Chiều cao: {cc} cm | Cân nặng: {cn} kg | BMI: {bmi} kg/m² ({eval_bmi})"
        pdf.add_txt(sh_line)

        if is_hau_phau:
            pdf.add_subsec("b. Vết mổ & Dẫn lưu:")
            pdf.add_txt(f"- Vết mổ: {payload.get('kham_vet_mo', '')}\n- Dẫn lưu: {payload.get('kham_dan_luu', '')}")
            pdf.add_subsec("c. Các cơ quan:")
        else:
            pdf.add_subsec("3. Thăm khám hiện tại - Các cơ quan:")

        organs = get_prioritized_organs(payload)

        for name, key in organs:
            pdf.add_subsec(f"- {name}:")
            pdf.add_txt(format_bullet_points(payload.get(key, "")))

        def sec_tom_tat(num_rom):
            pdf.add_sec(f"{num_rom}. TÓM TẮT BỆNH ÁN")
            pdf.add_txt(format_bullet_points(payload.get("tom_tat", "")))

        def sec_chan_doan_so_bo(num_sb, num_pb, num_bl):
            pdf.add_sec(f"{num_sb}. CHẨN ĐOÁN SƠ BỘ")
            pdf.add_txt(payload.get("chan_doan_so_bo", ""))
            pdf.add_sec(f"{num_pb}. CHẨN ĐOÁN PHÂN BIỆT")
            pdf.add_txt(payload.get("chan_doan_phan_biet", ""))
            if payload.get("bien_luan"):
                pdf.add_sec(f"{num_bl}. BIỆN LUẬN CHẨN ĐOÁN SƠ BỘ")
                pdf.add_txt(payload.get("bien_luan", ""))

        def sec_can_lam_sang(num_dx, num_co):
            pdf.add_sec(f"{num_dx}. ĐỀ XUẤT CẬN LÂM SÀNG")
            label_1 = "1. Đánh giá sau mổ / Biến chứng:" if is_hau_phau else "1. Phục vụ chẩn đoán xác định:"
            label_2 = "2. Theo dõi hồi phục & Chăm sóc:" if is_hau_phau else "2. Phục vụ điều trị:"
            pdf.add_subsec(label_1)
            pdf.add_txt(format_bullet_points(payload.get("cls_dx_xac_dinh", "")))
            pdf.add_subsec(label_2)
            pdf.add_txt(format_bullet_points(payload.get("cls_dx_dieu_tri", "")))
            pdf.add_subsec("3. Cận lâm sàng khác:")
            pdf.add_txt(format_bullet_points(payload.get("cls_dx_khac", "")))

            pdf.add_sec(f"{num_co}. CẬN LÂM SÀNG ĐÃ CÓ")
            cls_rows = []
            so_hang = int(payload.get("so_hang_cls", 3))
            for i in range(so_hang):
                kq = payload.get(f"cls_kq_{i}", "").strip()
                pg = payload.get(f"cls_pg_{i}", "").strip()
                img_list = payload.get(f"cls_img_list_{i}", [])
                if not img_list and payload.get(f"cls_img_b64_{i}"):
                    img_list = [payload.get(f"cls_img_b64_{i}")]
                
                if kq or pg or img_list:
                    cls_rows.append((kq, pg, img_list))
            if cls_rows:
                pdf.render_table_cls(cls_rows)
            else:
                pdf.add_txt("Chưa ghi nhận kết quả cận lâm sàng.")

        def sec_chan_doan_xac_dinh(num_xd, num_blxd):
            pdf.add_sec(f"{num_xd}. CHẨN ĐOÁN XÁC ĐỊNH")
            pdf.add_txt(format_bullet_points(payload.get("chan_doan_xac_dinh", "")))
            if payload.get("bien_luan_xac_dinh"):
                pdf.add_sec(f"{num_blxd}. BIỆN LUẬN CHẨN ĐOÁN XÁC ĐỊNH")
                pdf.add_txt(format_bullet_points(payload.get("bien_luan_xac_dinh", "")))

        if not is_hau_phau:
            sec_tom_tat("VI")
            sec_chan_doan_so_bo("VII", "VIII", "IX")
            sec_can_lam_sang("X", "XI")
            sec_chan_doan_xac_dinh("XII", "XIII")
            pdf.add_sec("XIV. ĐIỀU TRỊ")
            pdf.add_subsec("1. Mục tiêu điều trị:")
            pdf.add_txt(format_bullet_points(payload.get("dt_muc_tieu", "")))
            pdf.add_subsec("2. Điều trị cụ thể:")
            pdf.add_txt(format_bullet_points(payload.get("dt_cu_the", "")))
            pdf.add_subsec("3. Theo dõi sau điều trị:")
            pdf.add_txt(format_bullet_points(payload.get("dt_theo_doi", "")))
            pdf.add_sec("XV. TIÊN LƯỢNG")
            pdf.add_txt(format_bullet_points(payload.get("tien_luong", "")))
            pdf.add_sec("XVI. TƯ VẤN")
            pdf.add_txt(format_bullet_points(payload.get("tu_van", "")))
        else:
            sec_tom_tat("VI")
            sec_chan_doan_so_bo("VII", "VIII", "IX")
            sec_can_lam_sang("X", "XI")
            sec_chan_doan_xac_dinh("XII", "XIII")
            pdf.add_sec("XIV. ĐIỀU TRỊ HẬU PHẪU")
            pdf.add_subsec("1. Mục tiêu điều trị:")
            pdf.add_txt(format_bullet_points(payload.get("dt_muc_tieu", "")))
            pdf.add_subsec("2. Điều trị cụ thể:")
            pdf.add_txt(format_bullet_points(payload.get("dt_cu_the", "")))
            pdf.add_subsec("3. Theo dõi sau điều trị:")
            pdf.add_txt(format_bullet_points(payload.get("dt_theo_doi", "")))
            pdf.add_sec("XV. TIÊN LƯỢNG")
            pdf.add_txt(format_bullet_points(payload.get("tien_luong", "")))
            pdf.add_sec("XVI. TƯ VẤN")
            pdf.add_txt(format_bullet_points(payload.get("tu_van", "")))

        pdf_bytes = bytes(pdf.output())
        return Response(content=pdf_bytes, media_type="application/pdf", headers={"Content-Disposition": "inline; filename=benhan.pdf"})
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi kết xuất PDF: {str(e)}")

@app.post("/api/preview/docx", response_class=HTMLResponse)
async def preview_docx(data: dict):
    is_hau_phau = data.get("loai_benh_an") == "Hậu phẫu"

    def text(value: Any) -> str:
        return html.escape(str(value or ""))

    def field(label: str, value: Any) -> str:
        value_text = text(value)
        if not value_text.strip():
            return ""
        return f'<p><strong>{text(label)}:</strong> {value_text}</p>'

    def section(number: str, title: str, content: str) -> str:
        return f'<section><h2>{text(number)}. {text(title)}</h2>{content}</section>'

    content = section("I", "PHẦN HÀNH CHÍNH", "".join([
        field("1. Họ và tên", data.get("ho_ten")), field("2. Tuổi", data.get("tuoi")),
        field("3. Giới tính", data.get("gioi_tinh")), field("4. Dân tộc", data.get("dan_tok")),
        field("5. Nghề nghiệp", data.get("nghe_nghiep")), field("6. Khoa / Phòng điều trị", data.get("khoa_phong")),
        field("7. Ngày vào viện", data.get("ngay_vao_vien")), field("8. Ngày làm bệnh án", data.get("ngay_lam_benh_an")),
        field("9. Người làm bệnh án", data.get("sinh_vien")), field("10. Địa chỉ", data.get("dia_chi"))
    ]))
    content += section("II & III", "LÝ DO VÀO VIỆN VÀ BỆNH SỬ", "".join([
        field("1. Lý do vào viện", data.get("ly_do_vao_vien")),
        field("2. Tình trạng trước mổ", data.get("bs_truoc_mo")) if is_hau_phau else field("2. Bệnh sử", data.get("benh_su")),
        field("3. Diễn biến trong mổ", data.get("bs_trong_mo")) if is_hau_phau else "",
        field("4. Diễn biến sau mổ", data.get("bs_sau_mo")) if is_hau_phau else ""
    ]))
    content += section("IV", "TIỀN SỬ", "".join([
        field("1. Tiền sử nội khoa", data.get("ts_noi_khoa")),
        field("2. Tiền sử ngoại khoa & Dị ứng", data.get("ts_ngoai_khoa")),
        field("3. Tiền sử bản thân (Lối sống)", data.get("ts_loi_song")),
        field("4. Tiền sử gia đình", data.get("ts_gia_dinh"))
    ]))
    preview_organs = get_prioritized_organs(data)
    preview_exam_offset = 3 if is_hau_phau else 1
    content += section("V", "THĂM KHÁM LÂM SÀNG", "".join([
        field("1. Thời điểm khám", data.get("ngay_hau_phau")) if is_hau_phau else field("1. Thăm khám lúc vào viện", data.get("kham_vao_vien")),
        field("2. Vết mổ", data.get("kham_vet_mo")) if is_hau_phau else "",
        field("3. Dẫn lưu", data.get("kham_dan_luu")) if is_hau_phau else "",
        field("4. Khám toàn thân", data.get("kham_toan_than")) if is_hau_phau else field("2. Khám toàn thân", data.get("kham_toan_than")),
        field("5. Dấu hiệu sinh tồn", f"Mạch: {data.get('sh_mach', '')} ck/p | HA: {data.get('sh_ha', '')} mmHg | Nhiệt độ: {data.get('sh_nhiet_do', '')} °C | Nhịp thở: {data.get('sh_nhip_tho', '')} l/p | SpO2: {data.get('sh_spo2', '')}%") if is_hau_phau else field("3. Dấu hiệu sinh tồn", f"Mạch: {data.get('sh_mach', '')} ck/p | HA: {data.get('sh_ha', '')} mmHg | Nhiệt độ: {data.get('sh_nhiet_do', '')} °C | Nhịp thở: {data.get('sh_nhip_tho', '')} l/p | SpO2: {data.get('sh_spo2', '')}%"),
        *[field(f"{preview_exam_offset + 3 + index}. {label}", data.get(key)) for index, (label, key) in enumerate(preview_organs)]
    ]))
    content += section("VI", "TÓM TẮT BỆNH ÁN", field("1. Nội dung tóm tắt", data.get("tom_tat")))
    content += section("VII", "CHẨN ĐOÁN SƠ BỘ & PHÂN BIỆT", "".join([
        field("1. Chẩn đoán sơ bộ", data.get("chan_doan_so_bo")), field("2. Chẩn đoán phân biệt", data.get("chan_doan_phan_biet")), field("3. Biện luận sơ bộ", data.get("bien_luan"))
    ]))
    rows = []
    try:
        row_count = int(data.get("so_hang_cls", 0))
    except (TypeError, ValueError):
        row_count = 0
    for index in range(row_count):
        kq = data.get(f"cls_kq_{index}", "")
        pg = data.get(f"cls_pg_{index}", "")
        images = data.get(f"cls_img_list_{index}", []) or []
        if isinstance(images, str):
            images = [images]
        image_html = "".join(f'<img src="{text(image)}" alt="Ảnh cận lâm sàng">' for image in images if image)
        if kq or pg or image_html:
            rows.append(f"<tr><td>{text(kq)}{image_html}</td><td>{text(pg)}</td></tr>")
    table = '<table><thead><tr><th>KẾT QUẢ XÉT NGHIỆM & HÌNH ẢNH</th><th>PHIÊN GIẢI / BIỆN GIẢI</th></tr></thead><tbody>' + "".join(rows) + '</tbody></table>' if rows else "<p>Chưa ghi nhận kết quả cận lâm sàng.</p>"
    content += section("VIII", "CẬN LÂM SÀNG", "".join([
        field("1. CLS chẩn đoán", data.get("cls_dx_xac_dinh")), field("2. CLS điều trị", data.get("cls_dx_dieu_tri")), field("3. CLS khác", data.get("cls_dx_khac")),
        '<h3>Cận lâm sàng đã có</h3>', table
    ]))
    content += section("IX", "CHẨN ĐOÁN XÁC ĐỊNH", field("1. Chẩn đoán xác định", data.get("chan_doan_xac_dinh")) + field("2. Biện luận xác định", data.get("bien_luan_xac_dinh")))
    content += section("X", "ĐIỀU TRỊ & TIÊN LƯỢNG", "".join([
        field("1. Mục tiêu điều trị", data.get("dt_muc_tieu")), field("2. Điều trị cụ thể", data.get("dt_cu_the")),
        field("3. Theo dõi", data.get("dt_theo_doi")), field("4. Tiên lượng", data.get("tien_luong")), field("5. Tư vấn", data.get("tu_van"))
    ]))
    return f'''<!doctype html><html lang="vi"><head><meta charset="utf-8"><style>
        @page {{ size: A4; margin: 18mm; }}
        * {{ box-sizing: border-box; }} body {{ margin: 0; background: #e7e7e7; color: #222; font-family: "Times New Roman", serif; font-size: 12pt; line-height: 1.35; }}
        main {{ width: 210mm; max-width: 100%; min-height: 297mm; margin: 18px auto; padding: 18mm; background: #fff; box-shadow: 0 1px 8px #999; }}
        h1 {{ margin: 0 0 4px; color: #0a246a; text-align: center; font-size: 19pt; }} .subtitle {{ text-align: center; margin: 0 0 18px; font-style: italic; }}
        section {{ margin: 0 0 14px; break-inside: avoid; }} h2 {{ margin: 0 0 6px; padding: 5px 8px; color: #0a246a; background: #e1ebf5; border-bottom: 1px solid #9aaabd; font-size: 14pt; }}
        p {{ margin: 4px 0; white-space: pre-wrap; }} strong {{ color: #111; }} table {{ width: 100%; border-collapse: collapse; margin-top: 6px; table-layout: fixed; }} th, td {{ border: 1px solid #777; padding: 6px; vertical-align: top; white-space: pre-wrap; overflow-wrap: anywhere; }} th {{ background: #e6ebf5; font-size: 10pt; }} td {{ width: 50%; }} td img {{ display: block; max-width: 100%; max-height: 150px; margin: 6px 0; object-fit: contain; }}
        @media print {{ body {{ background: #fff; }} main {{ width: auto; min-height: auto; margin: 0; padding: 0; box-shadow: none; }} }}
    </style></head><body><main><h1>BỆNH ÁN LÂM SÀNG</h1><p class="subtitle">Loại hình: {text(data.get("loai_benh_an", "Nội khoa / Tiền phẫu"))}</p>{content}</main></body></html>'''


@app.post("/api/export/docx")
async def export_docx(data: dict):
    doc = Document()

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = DocxPt(12)
    normal_style.paragraph_format.space_after = DocxPt(4)
    normal_style.paragraph_format.line_spacing = 1.15

    # Cấu hình lề trang chuẩn văn bản y tế (1 inch ~ 2.54 cm)
    for section in doc.sections:
        section.top_margin = DocxInches(1)
        section.bottom_margin = DocxInches(1)
        section.left_margin = DocxInches(1)
        section.right_margin = DocxInches(1)

    # Tiêu đề bệnh án
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("BỆNH ÁN LÂM SÀNG")
    title_run.font.name = "Times New Roman"
    title_run.font.size = DocxPt(16)
    title_run.bold = True
    title_run.font.color.rgb = DocxRGBColor(10, 36, 106) # Xanh Win2K

    sub_title = doc.add_paragraph()
    sub_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    loai_ba = data.get("loai_benh_an", "Nội khoa / Tiền phẫu")
    r_sub = sub_title.add_run(f"Loại hình: {loai_ba}")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = DocxPt(12)
    r_sub.italic = True

    doc.add_paragraph() # Khoảng trống

    def add_section_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = DocxPt(8)
        p.paragraph_format.space_after = DocxPt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = DocxPt(13)
        run.bold = True
        run.font.color.rgb = DocxRGBColor(10, 36, 106)

    def add_field(label, val):
        if val and str(val).strip():
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = DocxPt(4)
            r_lbl = p.add_run(f"{label}: ")
            r_lbl.font.name = "Times New Roman"
            r_lbl.font.size = DocxPt(12)
            r_lbl.bold = True
            
            r_val = p.add_run(str(val))
            r_val.font.name = "Times New Roman"
            r_val.font.size = DocxPt(12)

    # 1. HÀNH CHÍNH
    add_section_heading("1. PHẦN HÀNH CHÍNH")
    add_field("1. Họ và tên", data.get("ho_ten"))
    add_field("2. Tuổi", data.get("tuoi"))
    add_field("3. Giới tính", data.get("gioi_tinh"))
    add_field("4. Dân tộc", data.get("dan_tok"))
    add_field("5. Nghề nghiệp", data.get("nghe_nghiep"))
    add_field("6. Khoa / Phòng điều trị", data.get("khoa_phong"))
    add_field("7. Ngày vào viện", data.get("ngay_vao_vien"))
    add_field("8. Ngày làm bệnh án", data.get("ngay_lam_benh_an"))
    add_field("9. Người làm bệnh án", data.get("sinh_vien"))
    add_field("10. Địa chỉ", data.get("dia_chi"))

    # 2. LÝ DO VÀO VIỆN
    add_section_heading("2. LÝ DO VÀO VIỆN")
    add_field("1. Lý do vào viện", data.get("ly_do_vao_vien"))

    # 3. BỆNH SỬ
    add_section_heading("3. BỆNH SỬ")
    if loai_ba == "Hậu phẫu":
        add_field("1. Tình trạng trước mổ", data.get("bs_truoc_mo"))
        add_field("2. Diễn biến trong mổ", data.get("bs_trong_mo"))
        add_field("3. Diễn biến sau mổ", data.get("bs_sau_mo"))
    else:
        add_field("1. Bệnh sử", data.get("benh_su"))

    # 4. TIỀN SỬ
    add_section_heading("4. TIỀN SỬ")
    add_field("1. Tiền sử nội khoa", data.get("ts_noi_khoa"))
    add_field("2. Tiền sử ngoại khoa & Dị ứng", data.get("ts_ngoai_khoa"))
    add_field("3. Tiền sử bản thân (Lối sống)", data.get("ts_loi_song"))
    add_field("4. Tiền sử gia đình", data.get("ts_gia_dinh"))

    # 5. THĂM KHÁM LÂM SÀNG
    add_section_heading("5. THĂM KHÁM LÂM SÀNG")
    if loai_ba == "Hậu phẫu":
        add_field("1. Thời điểm khám", data.get("ngay_hau_phau"))
        add_field("2. Vết mổ", data.get("kham_vet_mo"))
        add_field("3. Dẫn lưu", data.get("kham_dan_luu"))
        exam_offset = 3
    else:
        add_field("1. Thăm khám lúc vào viện", data.get("kham_vao_vien"))
        exam_offset = 1
    
    add_field(f"{exam_offset + 1}. Khám toàn thân", data.get("kham_toan_than"))
    sh_str = f"Mạch: {data.get('sh_mach', '')} ck/p | HA: {data.get('sh_ha', '')} mmHg | Nhiệt độ: {data.get('sh_nhiet_do', '')} °C | Nhịp thở: {data.get('sh_nhip_tho', '')} l/p | SpO2: {data.get('sh_spo2', '')}%"
    add_field(f"{exam_offset + 2}. Dấu hiệu sinh tồn", sh_str)
    
    organs = get_prioritized_organs(data)

    for number, (label, key) in enumerate(organs, start=exam_offset + 3):
        add_field(f"{number}. {label}", data.get(key))

    # 6. TÓM TẮT BỆNH ÁN
    add_section_heading("6. TÓM TẮT BỆNH ÁN")
    add_field("1. Nội dung tóm tắt", data.get("tom_tat"))

    # 7. CHẨN ĐOÁN SƠ BỘ & PHÂN BIỆT
    add_section_heading("7. CHẨN ĐOÁN SƠ BỘ & PHÂN BIỆT")
    add_field("1. Chẩn đoán sơ bộ", data.get("chan_doan_so_bo"))
    add_field("2. Chẩn đoán phân biệt", data.get("chan_doan_phan_biet"))
    add_field("3. Biện luận sơ bộ", data.get("bien_luan"))

    # 8. CẬN LÂM SÀNG
    add_section_heading("8. CẬN LÂM SÀNG")
    add_field("1. CLS chẩn đoán", data.get("cls_dx_xac_dinh"))
    add_field("2. CLS điều trị", data.get("cls_dx_dieu_tri"))
    add_field("3. CLS khác", data.get("cls_dx_khac"))

    # Bảng kết quả cận lâm sàng
    try:
        so_hang = max(0, int(data.get("so_hang_cls", 0) or 0))
    except (TypeError, ValueError):
        so_hang = 0
    if so_hang > 0:
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'KẾT QUẢ XÉT NGHIỆM & HÌNH ẢNH'
        hdr_cells[1].text = 'PHIÊN GIẢI / BIỆN GIẢI'
        for cell in hdr_cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = DocxPt(11)
                    r.bold = True

        for i in range(so_hang):
            kq = data.get(f"cls_kq_{i}", "")
            pg = data.get(f"cls_pg_{i}", "")
            img_list = data.get(f"cls_img_list_{i}", []) or []
            if isinstance(img_list, str):
                img_list = [img_list]

            if kq or pg or img_list:
                row_cells = table.add_row().cells
                p0 = row_cells[0].paragraphs[0]
                p0.paragraph_format.line_spacing = 1.1
                p0.add_run(str(kq or ""))
                
                # Chèn ảnh đính kèm (nếu có)
                for b64 in img_list:
                    try:
                        if "," in b64:
                            b64 = b64.split(",")[1]
                        img_bytes = io.BytesIO(base64.b64decode(b64))
                        row_cells[0].add_paragraph().add_run().add_picture(img_bytes, width=DocxInches(2.2))
                    except Exception:
                        pass

                p1 = row_cells[1].paragraphs[0]
                p1.paragraph_format.line_spacing = 1.1
                p1.add_run(str(pg or ""))

    # 9. CHẨN ĐOÁN XÁC ĐỊNH
    add_section_heading("9. CHẨN ĐOÁN XÁC ĐỊNH")
    add_field("1. Chẩn đoán xác định", data.get("chan_doan_xac_dinh"))
    add_field("2. Biện luận xác định", data.get("bien_luan_xac_dinh"))

    add_section_heading("10. ĐIỀU TRỊ, TIÊN LƯỢNG & TƯ VẤN")
    add_field("1. Mục tiêu điều trị", data.get("dt_muc_tieu"))
    add_field("2. Điều trị cụ thể", data.get("dt_cu_the"))
    add_field("3. Theo dõi", data.get("dt_theo_doi"))
    add_field("4. Tiên lượng", data.get("tien_luong"))
    add_field("5. Tư vấn", data.get("tu_van"))

    target_stream = io.BytesIO()
    doc.save(target_stream)
    target_stream.seek(0)

    ten_benh_nhan = str(data.get("ho_ten") or "Ho_So").strip()
    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", strip_accents(ten_benh_nhan)).strip("._") or "Ho_So"
    filename = f"Benh_An_{safe_name}.docx"
    encoded_filename = urllib.parse.quote(f"Benh_An_{ten_benh_nhan.replace(' ', '_')}.docx", safe="")
    return StreamingResponse(
        target_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}; filename*=UTF-8''{encoded_filename}"}
    )
